from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

try:
    import yaml
    from docx import Document
    from docx.document import Document as DocumentType
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph
except ImportError as exc:  # pragma: no cover
    raise SystemExit(f"python-docx and PyYAML are required: {exc}")

from python_engine import Issue, add_issue


DENY_PATTERNS = (
    ("tool_name", re.compile(r"\b(Codex|ChatGPT|Claude|GPT|Opus|Gemini)\b"), "Remove tool or model names."),
    ("source_trace", re.compile(r"文档由 AI|由模型|由智能体|由工具生成|本文由 AI"), "Remove generation traces."),
    ("local_path", re.compile(r"/Users/|file://|\.od/projects/|C:\\"), "Remove local paths."),
    ("old_structure", re.compile(r"基本信息表|字段｜内容|一、入口：|▌"), "Remove obsolete structure."),
)

INTERNAL_ID_RE = re.compile(r"\b[MS]\d{1,2}\b|screen\d+|page\d+", re.I)
TITLE_PREFIX_RE = re.compile(r"^\d+(?:\.\d+)*\.?\s+\S")
HEADING_RE = re.compile(r"^(Heading|标题)\s*([1-5])$")
SECTION_RE = re.compile(r"^(\d+)(?:\.|\s)")
RISK_KEYWORDS = ("资金", "奖励", "风控", "结算", "开奖", "赔率", "抽奖")
LITERAL_PATTERNS = ("<br>", "<br/>", r"\n")
CELL_BREAK_RE = re.compile(r"<(?:\w+:)?br(?:\s+[^>]*)?/?>")
TITLE_END_PUNCTUATION = ("。", "，", "、", "；", "！", "？")


@dataclass
class Block:
    kind: str
    obj: Paragraph | Table
    index: int
    section: str = ""
    heading_text: str = ""
    heading_level: int | None = None


def iter_block_items(parent: DocumentType | _Cell) -> Iterator[Paragraph | Table]:
    if isinstance(parent, DocumentType):
        parent_elm = parent.element.body
        parent_obj = parent
    else:
        parent_elm = parent._tc
        parent_obj = parent
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent_obj)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent_obj)


def heading_level_from_style(style_name: str) -> int | None:
    match = HEADING_RE.match(style_name)
    return int(match.group(2)) if match else None


def is_numbered_title(text: str) -> bool:
    stripped = text.strip()
    return bool(TITLE_PREFIX_RE.match(stripped)) and len(stripped) <= 60 and not stripped.rstrip().endswith(TITLE_END_PUNCTUATION)


def table_text(table: Table) -> str:
    parts: list[str] = []
    for row in table.rows:
        for cell in row.cells:
            parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def cell_image_count(cell: _Cell) -> int:
    return cell._tc.xml.count("<pic:pic")


def build_blocks(doc: DocumentType) -> list[Block]:
    blocks: list[Block] = []
    current_section = ""
    current_heading = ""
    current_level: int | None = None
    for index, item in enumerate(iter_block_items(doc), start=1):
        if isinstance(item, Paragraph):
            style = item.style.name if item.style is not None else ""
            text = item.text.strip()
            level = heading_level_from_style(style)
            if level and text:
                current_heading = text
                current_level = level
                section_match = SECTION_RE.match(text)
                if current_level == 1 and section_match:
                    current_section = section_match.group(1)
            blocks.append(Block("paragraph", item, index, current_section, current_heading, current_level))
        else:
            blocks.append(Block("table", item, index, current_section, current_heading, current_level))
    return blocks


def paragraph_texts(doc: DocumentType) -> list[tuple[str, str, str]]:
    texts: list[tuple[str, str, str]] = []
    p_index = 0
    t_index = 0
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            p_index += 1
            style_name = block.style.name if block.style is not None else ""
            context = "heading" if heading_level_from_style(style_name) is not None else "paragraph"
            texts.append((f"paragraph {p_index}", block.text, context))
        else:
            t_index += 1
            for r_idx, row in enumerate(block.rows, start=1):
                for c_idx, cell in enumerate(row.cells, start=1):
                    for p_idx, paragraph in enumerate(cell.paragraphs, start=1):
                        texts.append((f"table {t_index} r{r_idx}c{c_idx} p{p_idx}", paragraph.text, "table_cell"))
    return texts


def header_index(table: Table, header_name: str) -> int | None:
    if not table.rows:
        return None
    headers = [cell.text.strip() for cell in table.rows[0].cells]
    for idx, value in enumerate(headers):
        if value == header_name:
            return idx
    return None


def is_feature_table(table: Table) -> bool:
    if not table.rows:
        return False
    headers = {cell.text.strip() for cell in table.rows[0].cells}
    return {"模块", "展示", "需求描述"}.issubset(headers) or {"Module", "Evidence", "Requirement Details"}.issubset(headers)


def feature_tables(blocks: list[Block]) -> Iterator[tuple[int, str, Table]]:
    table_no = 0
    for block in blocks:
        if block.kind != "table":
            continue
        table_no += 1
        table = block.obj
        if isinstance(table, Table) and (block.section in ("4", "") or is_feature_table(table)) and is_feature_table(table):
            yield table_no, block.heading_text, table


def check_c1(doc: DocumentType, issues: list[Issue]) -> None:
    for location, text, context in paragraph_texts(doc):
        if not text:
            continue
        for _name, pattern, suggestion in DENY_PATTERNS:
            for match in pattern.finditer(text):
                add_issue(issues, "P01", "error", location, match.group(0), suggestion)
        if context in ("heading", "table_cell") or not (context == "paragraph" and len(text) > 30):
            for match in INTERNAL_ID_RE.finditer(text):
                add_issue(issues, "P01", "error", location, match.group(0), "Replace internal ids with business names.")


def check_c2(doc: DocumentType, issues: list[Issue], allow_loose_headings: bool) -> None:
    if allow_loose_headings:
        return
    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if not is_numbered_title(text):
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if heading_level_from_style(style_name) is None:
            add_issue(issues, "P02", "error", f"paragraph {idx}", f"{text} (style={style_name or 'None'})", "Numbered body headings must use Word Heading 1-5.")


def check_c3(doc: DocumentType, issues: list[Issue]) -> None:
    for t_idx, table in enumerate(doc.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            for c_idx, cell in enumerate(row.cells, start=1):
                if CELL_BREAK_RE.search(cell._tc.xml):
                    add_issue(issues, "P03", "error", f"table {t_idx} r{r_idx}c{c_idx}", "w:br", "Split table-cell content into paragraphs.")
                for p_idx, paragraph in enumerate(cell.paragraphs, start=1):
                    for literal in LITERAL_PATTERNS:
                        if literal in paragraph.text:
                            add_issue(issues, "P03", "error", f"table {t_idx} r{r_idx}c{c_idx} p{p_idx}", literal, "Do not write literal line breaks.")


def check_c4(doc: DocumentType, blocks: list[Block], issues: list[Issue], allow_missing_evidence: bool) -> None:
    if len(doc.inline_shapes) == 0:
        add_issue(issues, "P04", "error", "document", "0 images", "Embed real screenshots or mark Evidence as missing.")
    for table_no, heading, table in feature_tables(blocks):
        display_idx = header_index(table, "展示")
        module_idx = header_index(table, "模块")
        if display_idx is None:
            display_idx = header_index(table, "Evidence")
            module_idx = header_index(table, "Module")
        if display_idx is None:
            continue
        for r_idx, row in enumerate(table.rows[1:], start=2):
            display_cell = row.cells[display_idx]
            module_name = row.cells[module_idx].text.strip() if module_idx is not None else heading
            has_image = cell_image_count(display_cell) > 0
            has_fallback = "待设计补图" in display_cell.text or "missing" in display_cell.text.lower()
            if not has_image and not has_fallback:
                if allow_missing_evidence:
                    continue
                add_issue(issues, "P04", "error", f"table {table_no} r{r_idx}c{display_idx + 1}", module_name or heading or "empty evidence", "Evidence must contain an image or missing marker.")


def section_text(blocks: list[Block], section: str) -> str:
    parts: list[str] = []
    for block in blocks:
        if block.section != section:
            continue
        if isinstance(block.obj, Paragraph):
            parts.append(block.obj.text)
        elif isinstance(block.obj, Table):
            parts.append(table_text(block.obj))
    return "\n".join(parts)


def admin_config_rows(blocks: list[Block]) -> int:
    count = 0
    for block in blocks:
        if block.section != "5" or not isinstance(block.obj, Table):
            continue
        idx = header_index(block.obj, "配置项") or header_index(block.obj, "Config item")
        if idx is None:
            continue
        for row in block.obj.rows[1:]:
            value = row.cells[idx].text.strip()
            if value and value != "/":
                count += 1
    return count


def check_c5(blocks: list[Block], issues: list[Issue], strict: bool) -> None:
    feature_text = section_text(blocks, "4")
    mentions = feature_text.count("管理端可配置") + feature_text.lower().count("admin configurable")
    rows = admin_config_rows(blocks)
    if mentions > rows:
        add_issue(issues, "P05", "error" if strict else "warn", "section 4 -> section 5", f"mentions={mentions}, rows={rows}", "Configurable fields need admin mapping.")
    if rows > mentions and mentions > 0:
        add_issue(issues, "P05", "error" if strict else "warn", "section 5 -> section 4", f"rows={rows}, mentions={mentions}", "Admin rows may be over-expanded.")


def collect_heading_sections(blocks: list[Block]) -> list[tuple[Block, str]]:
    sections: list[tuple[Block, str]] = []
    for idx, block in enumerate(blocks):
        if not isinstance(block.obj, Paragraph):
            continue
        level = block.heading_level
        if not level:
            continue
        content: list[str] = []
        for next_block in blocks[idx + 1:]:
            if isinstance(next_block.obj, Paragraph) and next_block.heading_level and next_block.heading_level <= level:
                break
            if isinstance(next_block.obj, Paragraph):
                content.append(next_block.obj.text)
            elif isinstance(next_block.obj, Table):
                content.append(table_text(next_block.obj))
        sections.append((block, "\n".join(content)))
    return sections


def check_c6(blocks: list[Block], issues: list[Issue], strict: bool) -> None:
    if not strict:
        return
    for heading, content in collect_heading_sections(blocks):
        title = heading.obj.text.strip() if isinstance(heading.obj, Paragraph) else ""
        if heading.section != "4" or heading.heading_level not in (2, 3):
            continue
        if not any(keyword in title for keyword in RISK_KEYWORDS):
            continue
        missing = [label for label in ("AC", "异常路径", "owner") if label not in content]
        if missing:
            add_issue(issues, "P06", "error", title, "missing: " + ", ".join(missing), "High-risk modules require AC, exceptions, and owner.")


def chapter4_module_headings(blocks: list[Block]) -> list[str]:
    out: list[str] = []
    for block in blocks:
        if block.kind == "paragraph" and block.section == "4" and block.heading_level == 2 and isinstance(block.obj, Paragraph):
            text = block.obj.text.strip()
            if text:
                out.append(text)
    return out


def load_manifest(manifest_path: Path, issues: list[Issue]) -> dict | None:
    try:
        data = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
    except Exception as exc:
        add_issue(issues, "P08", "error", str(manifest_path), type(exc).__name__, "Cannot read manifest YAML.")
        return None
    if not isinstance(data, dict) or not isinstance(data.get("modules"), list):
        add_issue(issues, "P08", "error", "manifest.modules", str(data), "Manifest must be an object with modules array.")
        return None
    return data


def check_c8(blocks: list[Block], issues: list[Issue], strict: bool, manifest_path: Path | None) -> None:
    if manifest_path is None:
        return
    manifest = load_manifest(manifest_path, issues)
    if manifest is None:
        return
    prd_modules = set(chapter4_module_headings(blocks))
    manifest_names = {str(item.get("module_name", "")).strip() for item in manifest.get("modules", []) if isinstance(item, dict)}
    manifest_names.discard("")
    if prd_modules and prd_modules != manifest_names:
        add_issue(issues, "P08", "error", "manifest.modules", f"prd={sorted(prd_modules)}, manifest={sorted(manifest_names)}", "Manifest modules must match chapter 4 module headings.")
    for idx, item in enumerate([x for x in manifest.get("modules", []) if isinstance(x, dict)], start=1):
        source = item.get("source")
        name = str(item.get("module_name", "")).strip() or f"modules[{idx}]"
        if source == "missing":
            if not str(item.get("missing_reason", "")).strip() or not str(item.get("owner", "")).strip():
                add_issue(issues, "P08", "error" if strict else "warn", name, "missing metadata", "Missing screenshot entries require reason and owner.")
        else:
            try:
                dpr = float(item.get("device_pixel_ratio", 0))
            except (TypeError, ValueError):
                dpr = 0
            if dpr < 2.0:
                add_issue(issues, "P08", "error" if strict else "warn", name, f"device_pixel_ratio={item.get('device_pixel_ratio')!r}", "Non-missing entries require DPR >= 2.0.")


def validate_docx(path: Path, strict: bool, manifest_path: Path | None, options: dict) -> tuple[list[Issue], dict[str, int | str | bool]]:
    doc = Document(path)
    blocks = build_blocks(doc)
    issues: list[Issue] = []
    allow_loose_headings = bool(options.get("allow_loose_headings"))
    allow_missing_evidence = bool(options.get("allow_missing_evidence"))
    check_c1(doc, issues)
    check_c2(doc, issues, allow_loose_headings)
    check_c3(doc, issues)
    check_c4(doc, blocks, issues, allow_missing_evidence)
    check_c5(blocks, issues, strict)
    check_c6(blocks, issues, strict)
    check_c8(blocks, issues, strict, manifest_path)
    meta: dict[str, int | str | bool] = {
        "path": str(path),
        "strict": strict,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "issues": len(issues),
        "errors": sum(1 for issue in issues if issue.level == "error"),
        "warnings": sum(1 for issue in issues if issue.level == "warn"),
        "manifest": str(manifest_path) if manifest_path else "",
        "loose_headings": allow_loose_headings,
    }
    return issues, meta
